#!/usr/bin/env python3

import io
import datetime
import zipfile
from typing import BinaryIO

import openpyxl                 # type: ignore

from eatmud import ConciseRecord, DetailedRecord, get_irrs

__all__ = ('read_record', 'save_record')


ROW_BEG = 6
ROW_END = 10000
CELL_NAME = 'C1'
CELL_CODE = 'C2'
CELL_COMMENT = 'C3'
ROW_BEG_DETAILED = 9
ROW_END_DETAILED = 10000
CELL_NAME_DETAILED = 'C1'
CELL_CODE_DETAILED = 'C2'
CELL_COMMENT_DETAILED = 'C3'
COLUMNS = 'ABCDEFGHIJKLMNOPQRSTUVWXYZ'


class XLSXFormatError(Exception):
    pass


def read_record(filename: str, title: str,
                *args) -> ConciseRecord | DetailedRecord:
    """Read concise record or detailed record from given xlsx file.

    Parameters
    ----------
    filename: str
        Filename of the xlsx workbook.
    title: str
        Name of the worksheet.

    Returns
    -------
    record: Record or RecordDetailed
        The record in the worksheet.

    Note
    ----
    The library `openpyxl` may have some issues with workbook.close(),
    which may not actually close the file. So, the file is firstly read
    into a buffer, which is then passed to `load_workbook`. This can
    make sure the xlsx file is properly closed. This is inspired by
    `https://stackoverflow.com/questions/31416842/openpyxl-does-not-close-excel-workbook-in-read-only-mode`.
    """
    with open(filename, 'rb') as f:
        mem_file = io.BytesIO(f.read())
    wb = openpyxl.load_workbook(mem_file, read_only=True, data_only=True)
    try:
        ws = wb[title]
        t = ws['A1'].value
    finally:
        wb.close()
    if t == 'Record' or t == 'ConciseRecord':
        return read_concise_record(mem_file, title)
    if t == 'RecordDetailed':
        return read_detailed_record(mem_file, title)
    raise XLSXFormatError('sheet is not in standard format, use '
                          'read_record_compatible or '
                          'read_record_detailed_compatible instead')


def read_concise_record(filename: str | BinaryIO,
                        title: str) -> ConciseRecord:
    """Read concise record from a standard record file in xlsx form.

    Parameters
    ----------
    filename: str or stream
        Filename or file stream of the xlsx workbook.
    title: str
        Name of the worksheet.

    Returns
    -------
    record: Record
        The record in the worksheet.
    """
    wb = openpyxl.load_workbook(filename, read_only=True, data_only=True)
    try:
        ws = wb[title]
        name = ws[CELL_NAME].value
        code = ws[CELL_CODE].value
        comment = ws[CELL_COMMENT].value
        name = '' if name is None else name
        code = '' if code is None else code
        comment = '' if comment is None else comment
        record = ConciseRecord(name, code, comment)
        for row in ws.iter_rows(min_row=ROW_BEG, values_only=True):
            idx = row[0]
            # in case of empty row
            if idx is None:
                continue
            date, investment, present_value, comment = row[1:5]
            if comment is None:
                comment = ''
            record.append(datetime.date.fromisoformat(date),
                          float(investment),
                          float(present_value),
                          comment)
    finally:
        wb.close()
    return record


def read_detailed_record(filename: str | BinaryIO,
                         title: str) -> DetailedRecord:
    """Read detailed record from a standard detailed record file in xlsx form.

    Parameters
    ----------
    filename: str or stream
        Filename or file stream of the xlsx workbook.
    title: str
        Name of the worksheet.

    Returns
    -------
    record: DetailedRecord
        The detailed record in the worksheet.
    """
    wb = openpyxl.load_workbook(filename, read_only=True, data_only=True)
    try:
        ws = wb[title]
        name = ws[CELL_NAME_DETAILED].value
        code = ws[CELL_CODE_DETAILED].value
        comment = ws[CELL_COMMENT_DETAILED].value
        name = '' if name is None else name
        code = '' if code is None else code
        comment = '' if comment is None else comment
        record = DetailedRecord(name, code, comment)
        for row in ws.iter_rows(min_row=ROW_BEG_DETAILED, values_only=True):
            idx = row[0]
            # in case of empty row
            if idx is None:
                continue
            date, investment, nav, share, comment = row[1:6]
            if comment is None:
                comment = ''
            record.append(datetime.date.fromisoformat(date),
                          float(investment),
                          float(nav),
                          float(share),
                          comment)
    finally:
        wb.close()
    return record


def save_record(r: ConciseRecord | DetailedRecord,
                filename: str,
                title: str,
                ) -> None:
    """Save record or detailed record into a standard record file.

    Parameters
    ----------
    r: Record or RecordDetailed
        The record to be stored.
    filename: str
        Filename of the xlsx workbook.
    title: str
        Title of the worksheet.
    """

    if isinstance(r, ConciseRecord):
        save_concise_record(r, filename, title)
    elif isinstance(r, DetailedRecord):
        save_detailed_record(r, filename, title)
    else:
        raise TypeError('type of record not recognized')


def save_concise_record(r: ConciseRecord,
                        filename: str | BinaryIO,
                        title: str) -> None:
    """Save record into a standard record file in xlsx form.

    Parameters
    ----------
    r: Record
        The record to be stored.
    filename: str or stream
        Filename of file stream of the xlsx workbook.
    title: str
        Title of the worksheet.
    """
    # Prepare irr data.
    irr1 = get_irrs(r, 1)
    irr3 = get_irrs(r, 3)
    irr5 = get_irrs(r, 5)
    irrall = get_irrs(r, None)

    try:
        wb = openpyxl.load_workbook(filename)
    except (FileNotFoundError, zipfile.BadZipFile):
        wb = openpyxl.Workbook()
        wb[wb.sheetnames[0]].title = title

    try:
        try:
            ws = wb[title]
        except KeyError:
            ws = wb.create_sheet(title)

        ws.freeze_panes = 'A{}'.format(ROW_BEG)

        font = openpyxl.styles.Font(bold=True)
        fill_dark = openpyxl.styles.PatternFill(start_color='c0c0c0',
                                                fill_type='solid')
        fill_light = openpyxl.styles.PatternFill(start_color='d9d9d9',
                                                 fill_type='solid')
        fill_black = openpyxl.styles.PatternFill(start_color='000000',
                                                 fill_type='solid')

        # framework of the sheet
        ws.cell(1, 1, value='ConciseRecord').fill = fill_black
        values = ('名称', '代码', '备注',
                  '总投入', '当前资产', '总收益',
                  '序号', '时间', '变动', '当前', '备注',
                  '累计投入', '当前收益')
        positions = ('B1', 'B2', 'B3',
                     'F1', 'F2', 'F3',
                     'A5', 'B5', 'C5', 'D5', 'E5',
                     'F5', 'G5')
        for i, value in enumerate(values):
            cell = ws[positions[i]]
            cell.value = value
            cell.font = font
            cell.fill = fill_dark
        ws.merge_cells('C1:E1')
        ws.merge_cells('C2:E2')
        ws.merge_cells('C3:E3')
        ws.column_dimensions['A'].width = 5.0
        ws.column_dimensions['E'].width = 19.0
        for col_name in 'BCDFG':
            ws.column_dimensions[col_name].width = 12.0

        # Title for irr data.
        irr_names = ['1年年化', '3年年化', '5年年化', '总年化']
        for i, key in enumerate(irr_names):
            col = 8 + i
            ws.column_dimensions[COLUMNS[col-1]].width = 12.0
            cell = ws.cell(5, col, value=key)
            cell.font = font
            cell.fill = fill_dark

        # fill in basic information
        ws['C1'] = r.name
        ws['C2'] = r.code
        ws['C3'] = r.comment
        ws['G1'] = '=SUM(C{beg}:C{end})'.format(beg=ROW_BEG, end=ROW_END)
        ws['G3'] = '=G2-G1'
        ws['G1'].number_format = '0.00'
        ws['G3'].number_format = '0.00;[red]-0.00'
        for pos in 'G1', 'G3':
            ws[pos].fill = fill_light

        # fill in data
        for idx, rs in enumerate(r):
            date = f'{rs.date}'
            row = idx + ROW_BEG
            cells = ['{}{}'.format(col, row) for col in 'ABCDEFG']
            ws[cells[0]] = idx
            ws[cells[1]] = date
            ws[cells[1]].number_format = 'yyyy-mm-dd'
            ws[cells[2]] = rs.investment
            ws[cells[2]].number_format = '0.00'
            ws[cells[3]] = rs.present_value
            ws[cells[4]] = rs.comment
            # total investment
            if idx:
                ws[cells[5]] = '=F{i}+{inv}'.format(i=row-1, inv=cells[2])
            else:
                ws[cells[5]] = '={inv}'.format(inv=cells[2])
            # profit
            ws[cells[6]] = '={val}-{inv_t}'.format(val=cells[3],
                                                   inv_t=cells[5])
            ws[cells[6]].number_format = '0.00;[red]-0.00'
            for c in cells[5:7]:
                ws[c].fill = fill_light
            # irr_data
            for j, irr_list in enumerate([irr1, irr3, irr5, irrall]):
                col = 8 + j
                cell = ws.cell(row, col, value=irr_list[idx])
                cell.number_format = '0.00%'

        # alignment
        alignment_c = openpyxl.styles.Alignment(horizontal="center",
                                                vertical="center")
        alignment_l = openpyxl.styles.Alignment(horizontal="left",
                                                vertical="center")
        for rowi in ws.iter_rows():
            for cell in rowi:
                if cell.column == 5 and cell.row >= ROW_BEG:
                    cell.alignment = alignment_l
                else:
                    cell.alignment = alignment_c

        ws.views.sheetView[0].selection[0].sqref = cells[0]
        wb.save(filename)
    finally:
        wb.close()
        wb = None


def save_detailed_record(rd: DetailedRecord,
                         filename: str | BinaryIO,
                         title: str) -> None:
    """Save record into a standard detailed record file in xlsx form.

    Parameters
    ----------
    rd: DetailedRecord
        The detailed record to be stored.
    filename: str or stream
        Filename or file stream of the xlsx workbook.
    title: str
        Title of the worksheet.
    """
    # Prepare irr data.
    irr1 = get_irrs(rd, 1)
    irr3 = get_irrs(rd, 3)
    irr5 = get_irrs(rd, 5)
    irrall = get_irrs(rd, None)

    try:
        wb = openpyxl.load_workbook(filename)
    except (FileNotFoundError, zipfile.BadZipFile):
        wb = openpyxl.Workbook()
        wb[wb.sheetnames[0]].title = title
    try:
        try:
            ws = wb[title]
        except KeyError:
            ws = wb.create_sheet(title)

        ws.freeze_panes = 'A{}'.format(ROW_BEG_DETAILED)

        font = openpyxl.styles.Font(bold=True)
        fill_dark = openpyxl.styles.PatternFill(start_color='c0c0c0',
                                                fill_type='solid')
        fill_light = openpyxl.styles.PatternFill(start_color='d9d9d9',
                                                 fill_type='solid')
        fill_black = openpyxl.styles.PatternFill(start_color='000000',
                                                 fill_type='solid')

        # framework of the sheet
        ws.cell(1, 1, value='RecordDetailed').fill = fill_black
        values = ('名称', '代码', '备注',
                  '总投入', '总份额', '总费用', '当前净值', '总资产', '总收益',
                  '序号', '时间', '投入费用', '交易净值', '交易份额', '备注',
                  '交易费用', '累计投入', '累计份额', '当前市值', '当前收益')
        positions = ('B1', 'B2', 'B3',
                     'B4', 'B5', 'B6', 'D4', 'D5', 'D6',
                     'A8', 'B8', 'C8', 'D8', 'E8', 'F8',
                     'G8', 'H8', 'I8', 'J8', 'K8')
        for i, value in enumerate(values):
            cell = ws[positions[i]]
            cell.value = value
            cell.font = font
            cell.fill = fill_dark
        ws.merge_cells('C1:E1')
        ws.merge_cells('C2:E2')
        ws.merge_cells('C3:E3')
        ws.column_dimensions['A'].width = 5.0
        ws.column_dimensions['F'].width = 19.0
        for col in 'BCDEGHIJK':
            ws.column_dimensions[col].width = 12.0

        # fill in basic information
        ws['C1'] = rd.name
        ws['C2'] = rd.code
        ws['C3'] = rd.comment
        ws['C4'] = '=SUM(C{beg}:C{end})'.format(beg=ROW_BEG_DETAILED,
                                                end=ROW_END_DETAILED)
        ws['C5'] = '=SUM(E{beg}:E{end})'.format(beg=ROW_BEG_DETAILED,
                                                end=ROW_END_DETAILED)
        ws['C6'] = '=SUM(G{beg}:G{end})'.format(beg=ROW_BEG_DETAILED,
                                                end=ROW_END_DETAILED)
        ws['E5'] = '=C5*E4'
        ws['E6'] = '=E5-C4'
        ws['C6'].number_format = '0.00'
        ws['E5'].number_format = '0.00'
        ws['E6'].number_format = '0.00;[red]-0.00'
        for pos in 'C4', 'C5', 'C6', 'E5', 'E6':
            ws[pos].fill = fill_light

        # Title for irr data.
        irr_names = ['1年年化', '3年年化', '5年年化', '总年化']
        for i, key in enumerate(irr_names):
            col_num = 12 + i
            ws.column_dimensions[COLUMNS[col_num-1]].width = 12.0
            cell = ws.cell(8, col_num, value=key)
            cell.font = font
            cell.fill = fill_dark

        # fill in data
        for idx, rs in enumerate(rd):
            date = f'{rs.date}'
            row = idx + ROW_BEG_DETAILED
            cells = ['{}{}'.format(col, row) for col in 'ABCDEFGHIJK']
            ws[cells[0]] = idx
            ws[cells[1]] = date
            ws[cells[1]].number_format = 'yyyy-mm-dd'
            ws[cells[2]] = rs.investment
            ws[cells[2]].number_format = '0.00'
            ws[cells[3]] = rs.nav
            ws[cells[4]] = rs.share
            ws[cells[4]].number_format = '0.0000'
            ws[cells[5]] = rs.comment
            # fee
            ws[cells[6]] = '={inv}-{nav}*{share}'.format(inv=cells[2],
                                                         nav=cells[3],
                                                         share=cells[4])
            ws[cells[6]].number_format = '0.0000'
            # total investment
            if idx:
                ws[cells[7]] = '=H{i}+{inv}'.format(i=row-1, inv=cells[2])
            else:
                ws[cells[7]] = '={inv}'.format(inv=cells[2])
            # total share
            if idx:
                ws[cells[8]] = '=I{i}+{share}'.format(i=row-1, share=cells[4])
            else:
                ws[cells[8]] = '={share}'.format(share=cells[4])
            # present value
            ws[cells[9]] = '={ts}*{nav}'.format(ts=cells[8], nav=cells[3])
            ws[cells[9]].number_format = '0.00'
            # profit
            ws[cells[10]] = '={val}-{inv_t}'.format(val=cells[9],
                                                    inv_t=cells[7])
            ws[cells[10]].number_format = '0.00;[red]-0.00'
            for c in cells[6:11]:
                ws[c].fill = fill_light

            # irr_data
            for j, irr_list in enumerate([irr1, irr3, irr5, irrall]):
                col_num = 12 + j
                cell = ws.cell(row, col_num, value=irr_list[idx])
                cell.number_format = '0.00%'

        # alignment
        alignment_c = openpyxl.styles.Alignment(horizontal="center",
                                                vertical="center")
        alignment_l = openpyxl.styles.Alignment(horizontal="left",
                                                vertical="center")
        for rowi in ws.iter_rows():
            for cell in rowi:
                if cell.column == 6 and cell.row >= ROW_BEG_DETAILED:
                    cell.alignment = alignment_l
                else:
                    cell.alignment = alignment_c

        ws.views.sheetView[0].selection[0].sqref = cells[0]
        wb.save(filename)
    finally:
        wb.close()
        wb = None
